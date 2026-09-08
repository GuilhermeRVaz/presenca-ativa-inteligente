"""
generate_official_ai_report.py — Gerador Oficial de Relatórios com IA (OpenAI + python-docx)

Gera os relatórios oficiais padronizados nos formatos Markdown (.md) e Word (.docx)
seguindo estritamente a identidade visual e o protocolo pedagógico da
E.E. Profa. Décia L. M. dos Santos.
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

load_dotenv(ROOT / ".env")

from app.core.config import settings
from app.infrastructure.supabase.repositories import SupabaseRepository

SCHOOL_NAME = "E.E. Profa. Décia L. M. dos Santos"


def fetch_campaign_data(absence_days: str | None = None, campaign_id: str | None = None) -> dict[str, Any]:
    """
    Busca dados completos das campanhas, mensagens e respostas do Supabase.
    """
    repo = SupabaseRepository()
    client = repo.client.schema("busca_ativa_v2")

    if campaign_id:
        camps_res = client.table("campaigns").select("*").in_("id", campaign_id.split(",")).execute()
    elif absence_days:
        camps_res = client.table("campaigns").select("*").eq("absence_days", absence_days).execute()
    else:
        # Pega a última campanha
        latest = client.table("campaigns").select("*").order("created_at", desc=True).limit(1).execute()
        if not latest.data:
            raise ValueError("Nenhuma campanha encontrada no banco.")
        absence_days = latest.data[0]["absence_days"]
        camps_res = client.table("campaigns").select("*").eq("absence_days", absence_days).execute()

    campaigns = camps_res.data or []
    if not campaigns:
        raise ValueError(f"Nenhuma campanha encontrada para {absence_days or campaign_id}.")

    camp_ids = [c["id"] for c in campaigns]
    date_str = campaigns[0].get("absence_days", absence_days or "N/A")

    # Mensagens
    msgs_res = (
        client.table("messages")
        .select("id, campaign_id, status, last_error, wa_jid, student_id, students(id, name, ra, class_name)")
        .in_("campaign_id", camp_ids)
        .execute()
    )
    messages = msgs_res.data or []

    # Respostas
    resp_res = (
        client.table("responses")
        .select("id, campaign_id, student_id, body, reason, received_at, created_at, students(name, ra, class_name)")
        .in_("campaign_id", camp_ids)
        .execute()
    )
    responses = resp_res.data or []

    return {
        "absence_days": date_str,
        "campaigns": campaigns,
        "messages": messages,
        "responses": responses
    }


def generate_markdown_with_ai(data: dict[str, Any]) -> str:
    """
    Chama a API da OpenAI para estruturar e sintetizar o relatório oficial em Markdown.
    """
    from openai import OpenAI

    api_key = settings.openai_api_key or os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise ValueError("OPENAI_API_KEY não configurada.")

    client = OpenAI(api_key=api_key)

    system_prompt = f"""Você é o especialista auditores do sistema Presença Ativa Inteligente (PAI).
Sua missão é gerar um RELATÓRIO OFICIAL DE OPERAÇÃO — BUSCA ATIVA para a unidade escolar '{SCHOOL_NAME}'.

DIRETRIZES OBRIGATÓRIAS:
1. O nome da escola DEVE SER EXATAMENTE: '{SCHOOL_NAME}'.
2. O tom deve ser altamente profissional, rigoroso, pedagógico e auditado.
3. Estrutura das Seções do Relatório em Markdown:
   - # Relatório Oficial de Operação — Busca Ativa (DD/MM/YYYY)
   - ## Versão Auditada e Consolidada | {SCHOOL_NAME} — Presença Ativa Inteligente (PAI)
   - Bloco de Contexto Operacional (com alertas GitHub em blockquote > ℹ️)
   - ## 📊 1. Resumo Executivo e Métricas da Campanha
   - ## ❌ 2. Alunos que Não Receberam Mensagens por Falha Técnica de Disparo (se houver)
   - ## 📋 3. Detalhamento por Turma — Todos os Alunos Monitorados (agrupados por turma com tabelas contendo: Nº, Aluno, RA, Protocolo, Status, Resposta da Família, Categoria)
   - ## 🔔 4. Casos que Requerem Atenção Especial e Alertas Prioritários (Tabela com destaques de Bullying, Ansiedade, Lesão na Escola, Gastroenterite, Dentista, Presenças Confirmadas, etc.)
   - ## 📌 5. Encaminhamentos e Orientações Pedagógicas
   - Rodapé: *Documento gerado em: [Data] | Sistema: Presença Ativa Inteligente (PAI) | {SCHOOL_NAME}*

Retorne APENAS o conteúdo em Markdown limpo."""

    user_content = f"""Gere o relatório oficial com base nos seguintes dados reais auditados:

Data das Faltas: {data['absence_days']}
Campanhas Executadas: {json.dumps(data['campaigns'], ensure_ascii=False, default=str)}
Resumo de Mensagens: {len(data['messages'])} registros.
Lista de Mensagens: {json.dumps(data['messages'][:100], ensure_ascii=False, default=str)}
Respostas Coletadas: {json.dumps(data['responses'], ensure_ascii=False, default=str)}
"""

    response = client.chat.completions.create(
        model="gpt-4o",
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_content}
        ],
        temperature=0.2
    )

    return response.choices[0].message.content or ""


def create_word_docx_from_markdown(md_content: str, output_filepath: str):
    """
    Converte ou gera o documento Word (.docx) com formatação executiva usando python-docx.
    """
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()
    section = doc.sections[0]
    section.page_width  = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin   = Cm(2.0)
    section.right_margin  = Cm(2.0)
    section.top_margin    = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    COR_AZUL_ESCURO = RGBColor(0x1A, 0x37, 0x6C)
    COR_AZUL_MEDIO  = RGBColor(0x2E, 0x5F, 0xA3)
    COR_BRANCO      = RGBColor(0xFF, 0xFF, 0xFF)

    # Título Principal
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Relatório Oficial de Operação — Busca Ativa')
    r.bold = True; r.font.size = Pt(18); r.font.color.rgb = COR_AZUL_ESCURO

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(f'Relatório Consolidado por IA | {SCHOOL_NAME} — Presença Ativa Inteligente (PAI)')
    r2.font.size = Pt(10.5); r2.font.color.rgb = COR_AZUL_MEDIO; r2.bold = True

    p3 = doc.add_paragraph('─' * 90)
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.runs[0].font.color.rgb = COR_AZUL_MEDIO

    # Processa cada linha do markdown simples
    lines = md_content.split('\n')
    table_data = []

    for line in lines:
        stripped = line.strip()
        if stripped.startswith('# '):
            pass
        elif stripped.startswith('## '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            r = h.add_run(stripped[3:])
            r.bold = True; r.font.size = Pt(13); r.font.color.rgb = COR_AZUL_MEDIO
        elif stripped.startswith('### '):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            r = h.add_run(stripped[4:])
            r.bold = True; r.font.size = Pt(11); r.font.color.rgb = COR_AZUL_ESCURO
        elif stripped.startswith('> '):
            box = doc.add_paragraph()
            box.paragraph_format.left_indent = Cm(0.5)
            r = box.add_run(stripped[2:])
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x30, 0x30, 0x30)
        elif stripped.startswith('|'):
            # Tabela markdown
            cols = [c.strip() for c in stripped.split('|')[1:-1]]
            if all(set(c) <= set('-: ') for c in cols):
                continue  # Separador de cabeçalho
            table_data.append(cols)
        else:
            if table_data:
                # Renderiza tabela acumulada
                headers = table_data[0]
                rows = table_data[1:]
                if headers and rows:
                    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
                    t.style = 'Table Grid'
                    t.alignment = WD_TABLE_ALIGNMENT.CENTER
                    hdr = t.rows[0]
                    for i, h_text in enumerate(headers):
                        cell = hdr.cells[i]
                        cell.text = h_text
                        tcPr = cell._tc.get_or_add_tcPr()
                        shd = OxmlElement('w:shd')
                        shd.set(qn('w:val'), 'clear')
                        shd.set(qn('w:color'), 'auto')
                        shd.set(qn('w:fill'), '1A376C')
                        tcPr.append(shd)
                        p = cell.paragraphs[0]
                        p.runs[0].font.color.rgb = COR_BRANCO
                        p.runs[0].bold = True
                        p.runs[0].font.size = Pt(8.5)
                    for ri, row_data in enumerate(rows):
                        r_elem = t.rows[ri + 1]
                        for ci, val in enumerate(row_data[:len(headers)]):
                            r_elem.cells[ci].text = val
                            p = r_elem.cells[ci].paragraphs[0]
                            p.runs[0].font.size = Pt(8)
                doc.add_paragraph()
                table_data = []

            if stripped:
                p = doc.add_paragraph(stripped)
                p.paragraph_format.space_after = Pt(3)

    doc.save(output_filepath)
    print(f"Documento Word salvo com sucesso em: {output_filepath}")


def generate_official_report(absence_days: str | None = None, campaign_id: str | None = None) -> dict[str, str]:
    """
    Função principal que gera os relatórios em Markdown e Word (.docx).
    """
    data = fetch_campaign_data(absence_days=absence_days, campaign_id=campaign_id)
    date_formatted = data['absence_days'].replace("/", "_")

    md_content = generate_markdown_with_ai(data)

    rel_dir = ROOT / "relatorios"
    rel_dir.mkdir(exist_ok=True)

    md_file = rel_dir / f"RELATORIO_OFICIAL_BUSCA_ATIVA_DIA_{date_formatted}.md"
    docx_file = rel_dir / f"RELATORIO_OFICIAL_BUSCA_ATIVA_DIA_{date_formatted}.docx"

    md_file.write_text(md_content, encoding="utf-8")
    print(f"Relatório Markdown salvo com sucesso em: {md_file}")

    create_word_docx_from_markdown(md_content, str(docx_file))

    return {
        "markdown_path": str(md_file),
        "docx_path": str(docx_file),
        "markdown_content": md_content
    }


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Gerador Oficial de Relatórios PAI com IA")
    parser.add_argument("--day", type=str, help="Dia das faltas (ex: 28/08/2026 ou 28)")
    parser.add_argument("--campaign-id", type=str, help="UUID da campanha")
    args = parser.parse_args()

    day_param = args.day
    if day_param and "/" not in day_param:
        now = datetime.now()
        day_param = f"{int(day_param):02d}/{now.month:02d}/{now.year}"

    generate_official_report(absence_days=day_param, campaign_id=args.campaign_id)
