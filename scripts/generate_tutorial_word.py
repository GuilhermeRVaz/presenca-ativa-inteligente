import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D0D7DE", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def add_callout_box(doc, text_list, title="NOTA IMPORTANTE", border_color="1F4E78", bg_color="F0F4F8"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    r_title = p.add_run(f"📌 {title}")
    r_title.bold = True
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(11)
    r_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
    
    for line in text_list:
        p_line = cell.add_paragraph()
        p_line.paragraph_format.space_before = Pt(2)
        p_line.paragraph_format.space_after = Pt(2)
        r = p_line.add_run(line)
        r.font.name = 'Calibri'
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_heading_styled(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    
    if level == 1:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(15)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78) # Navy Blue
    elif level == 2:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6) # Steel Blue
    elif level == 3:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
    return p

def build_tutorial_docx():
    doc = Document()
    
    # Configuração de Margens (2cm)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
        # Cabeçalho da página
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("Presença Ativa Inteligente (PAI) — Manual de Replicação")
        hr.font.name = 'Calibri'
        hr.font.size = Pt(8.5)
        hr.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    # Estilo Normal
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x26, 0x26, 0x26)

    # -------------------------------------------------------------
    # CAPA / TÍTULO PRINCIPAL
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("MANUAL OFICIAL DE REPLICAÇÃO DO SISTEMA")
    run_title.bold = True
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Presença Ativa Inteligente (PAI) — Guia Passo a Passo de Implantação em Novas Escolas")
    run_sub.font.size = Pt(12.5)
    run_sub.font.color.rgb = RGBColor(0x59, 0x59, 0x59)

    add_callout_box(
        doc,
        [
            "Este documento foi elaborado para permitir que qualquer gestor escolar, técnico de TI ou usuário iniciante implante o ecossistema completo do PAI.",
            "O sistema pode ser construído, configurado e customizado de forma rápida e autônoma utilizando a IDE Antigravity com recursos de IA e Vibecoding.",
            "Contém checklist de Pen Drive, obtenção de chaves (.env), banco Supabase, WhatsApp, Docker e Prompts Prontos."
        ],
        title="RESUMO EXECUTIVO DO GUIA DE REPLICAÇÃO",
        border_color="1F4E78",
        bg_color="F0F4F8"
    )

    # -------------------------------------------------------------
    # SEÇÃO 1: VISÃO GERAL DA ARQUITETURA
    # -------------------------------------------------------------
    add_heading_styled(doc, "1. Visão Geral e Arquitetura do Sistema", level=1)
    
    p = doc.add_paragraph()
    p.add_run(
        "O Presença Ativa Inteligente (PAI) é uma plataforma de combate à evasão escolar e busca ativa automatizada. "
        "O fluxo de comunicação integra WhatsApp oficial, inteligência artificial generativa, banco de dados seguro em nuvem e um painel web simplificado para a equipe pedagógica."
    )

    tbl_arch = doc.add_table(rows=6, cols=3)
    tbl_arch.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_arch.autofit = False
    set_table_borders(tbl_arch)

    headers = ["Componente", "Tecnologia / Porta", "Função no Ecossistema"]
    col_widths = [Inches(1.8), Inches(1.8), Inches(3.2)]

    for i, h in enumerate(headers):
        cell = tbl_arch.cell(0, i)
        cell.width = col_widths[i]
        set_cell_background(cell, "1F4E78")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)

    arch_data = [
        ("Painel de Gestão", "Streamlit (Porta 8501)", "Interface visual para secretaria: envio de faltas, justificativas e relatórios."),
        ("Backend API", "FastAPI (Porta 8000)", "Motor de regras de negócio, validação de contatos e endpoints REST seguros."),
        ("Orquestrador", "n8n (Porta 5678)", "Gerencia os fluxos de webhooks, roteamento e automação de disparos."),
        ("Gateway WhatsApp", "Evolution API / Z-API", "Conecta o chip oficial de WhatsApp da escola ao sistema via API."),
        ("Inteligência Artificial", "OpenAI / Gemini LLM", "Classificação de intenções, triagem de justificativas e FAQ escolar.")
    ]

    for row_idx, data in enumerate(arch_data, start=1):
        bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(data):
            cell = tbl_arch.cell(row_idx, col_idx)
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if col_idx == 0:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SEÇÃO 2: CHECKLIST PEN DRIVE
    # -------------------------------------------------------------
    add_heading_styled(doc, "2. O Que Copiar no Pen Drive (Kit de Implantação)", level=1)
    
    p = doc.add_paragraph()
    p.add_run(
        "Ao preparar o Pen Drive para levar à nova escola, copie apenas os arquivos estruturais e limpos. "
        "Nunca copie o arquivo .env com senhas da escola anterior ou pastas temporárias."
    )

    tbl_pen = doc.add_table(rows=10, cols=3)
    tbl_pen.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pen.autofit = False
    set_table_borders(tbl_pen)

    pen_headers = ["Item / Pasta", "Status", "Descrição e Finalidade"]
    pen_widths = [Inches(2.0), Inches(1.2), Inches(3.6)]

    for i, h in enumerate(pen_headers):
        cell = tbl_pen.cell(0, i)
        cell.width = pen_widths[i]
        set_cell_background(cell, "2E75B6")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        r.font.size = Pt(9.5)

    pen_data = [
        ("app/ e pages/", "✅ Copiar", "Código fonte da API FastAPI e páginas do Painel."),
        ("migrations/", "✅ Copiar", "Scripts SQL de criação do banco de dados (pasta versions/)."),
        ("scripts/", "✅ Copiar", "Utilitários de importação de alunos e relatórios Word."),
        ("docker-compose.yml e Dockerfile", "✅ Copiar", "Arquivos de inicialização de containers."),
        ("requirements.txt e painel.py", "✅ Copiar", "Dependências Python e tela principal Streamlit."),
        ("workflow_triagem_final.json", "✅ Copiar", "Fluxo n8n pronto para importação."),
        (".env.example", "✅ Copiar", "Modelo de variáveis de configuração em branco."),
        ("TUTORIAL e Manual .docx", "✅ Copiar", "Documentação completa de implantação."),
        (".env, .venv, .git, dados antigos", "❌ NÃO COPIAR", "Segurança: Nunca transporte credenciais ou dados de alunos antigos.")
    ]

    for row_idx, data in enumerate(pen_data, start=1):
        bg = "F9FAFB" if row_idx % 2 == 1 else "FFFFFF"
        if "NÃO COPIAR" in data[1]:
            bg = "FDEDEC"
        for col_idx, text in enumerate(data):
            cell = tbl_pen.cell(row_idx, col_idx)
            cell.width = pen_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(9.5)
            if col_idx == 0 or "NÃO COPIAR" in text:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # SEÇÃO 3: ONDE E COMO CONSEGUIR TODAS AS CHAVES (.ENV)
    # -------------------------------------------------------------
    add_heading_styled(doc, "3. Como Obter Todas as Chaves de Configuração (.env)", level=1)
    
    p = doc.add_paragraph()
    p.add_run("Renomeie o arquivo ")
    r_bold = p.add_run(".env.example")
    r_bold.bold = True
    p.add_run(" para ")
    r_bold2 = p.add_run(".env")
    r_bold2.bold = True
    p.add_run(" e preencha as variáveis obtidas nos 4 serviços externos:")

    add_heading_styled(doc, "A. Supabase (Banco de Dados em Nuvem)", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Acesse ")
    p.add_run("https://supabase.com").bold = True
    p.add_run(" e crie um projeto gratuito com o nome da escola.\n")
    p.add_run("2. Em ")
    p.add_run("Project Settings > API").bold = True
    p.add_run(", copie a ")
    p.add_run("Project URL").bold = True
    p.add_run(" (vai em SUPABASE_URL) e a chave ")
    p.add_run("anon public").bold = True
    p.add_run(" (vai em SUPABASE_KEY).\n")
    p.add_run("3. Abra o ")
    p.add_run("SQL Editor").bold = True
    p.add_run(" no Supabase e execute os 7 arquivos da pasta migrations/versions/ em ordem sequencial (0001 a 0007).")

    add_heading_styled(doc, "B. Evolution API (WhatsApp da Escola)", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Acesse o painel da Evolution API e crie uma nova instância (ex: escola_modelo).\n")
    p.add_run("2. Anote a URL da API (EVOLUTION_API_URL) e a Chave de Acesso (EVOLUTION_API_KEY).\n")
    p.add_run("3. Clique em Gerar QR Code e conecte o WhatsApp oficial da escola no celular da secretaria.\n")
    p.add_run("4. Aponte o Webhook para o endereço do seu n8n (/webhook/triagem).")

    add_heading_styled(doc, "C. OpenAI / Gemini (Inteligência Artificial)", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Acesse ")
    p.add_run("https://platform.openai.com").bold = True
    p.add_run(" e gere uma nova API Key (sk-proj-...).\n")
    p.add_run("2. Cole em OPENAI_API_KEY no .env. Essa chave permite que a IA classifique justificativas e responda dúvidas de pais.")

    add_heading_styled(doc, "D. n8n (Orquestrador de Fluxos)", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Com o sistema rodando, acesse ")
    p.add_run("http://localhost:5678").bold = True
    p.add_run(" no navegador.\n")
    p.add_run("2. Vá em Workflows > Import from File e selecione ")
    p.add_run("workflow_triagem_final.json").bold = True
    p.add_run(".\n3. Ative o fluxo no botão superior direito (Active).")

    # -------------------------------------------------------------
    # SEÇÃO 4: CONSTRUÇÃO COM IDE ANTIGRAVITY & VIBECODING
    # -------------------------------------------------------------
    add_heading_styled(doc, "4. Construção, Vibecoding, MCPs e Skills na IDE Antigravity", level=1)
    
    p = doc.add_paragraph()
    p.add_run(
        "A IDE Antigravity permite implantar e customizar o sistema através de comandos em linguagem natural (Vibecoding). "
        "Ao conectar os servidores MCP e ativar as Skills especializadas, o assistente de IA ganha autonomia total para manipular banco de dados, fluxos do n8n e criar relatórios."
    )

    add_callout_box(
        doc,
        [
            "MCP (Model Context Protocol): Conecta a IA diretamente ao Supabase e n8n para ler e executar comandos.",
            "Skills: Conjuntos de instruções de especialistas (n8n-mcp-tools-expert, supabase-automation, fastapi-pro) que garantem código limpo e sem bugs.",
            "Vibecoding: Você apenas digita o que precisa (ex: 'Cadastre os alunos da turma 3B') e a Antigravity executa!"
        ],
        title="O QUE SÃO MCPs E SKILLS NO ANTIGRAVITY?",
        border_color="C55A11",
        bg_color="FEF9E7"
    )

    add_heading_styled(doc, "Prompts Mestres Prontos ('Plug and Play') para Colar na Antigravity", level=2)

    prompts = [
        ("Prompt 1: Validação Geral de Ambiente", 
         "\"Olá Antigravity! Estou implantando o sistema PAI para a [Nome da Escola]. Acabei de preencher o arquivo .env. Por favor, verifique a conexão com o Supabase, teste as variáveis de ambiente e execute os testes unitários (pytest) para garantir que tudo está pronto.\""),
        ("Prompt 2: Verificação do Banco de Dados",
         "\"Antigravity, verifique o status das tabelas no Supabase. Compare com os arquivos da pasta migrations/versions/ e garanta que todas as tabelas e permissões estão 100% criadas e prontas.\""),
        ("Prompt 3: Importação de Alunos e Turmas",
         "\"Antigravity, aqui está a lista de alunos da escola [anexar CSV ou colar nomes/telefones]. Crie e execute um script seguro em scripts/ para importar todos os alunos no Supabase associados à nossa escola.\""),
        ("Prompt 4: Ajuste das Regras Escolares na IA (FAQ/RAG)",
         "\"Antigravity, as regras da nossa escola são: entrada às 07h00, portão fecha às 07h15, uniforme obrigatório e reuniões bimestrais. Atualize o prompt do assistente de triagem para responder essas dúvidas aos pais.\""),
        ("Prompt 5: Inicialização e Teste de Saúde",
         "\"Antigravity, inicie todos os containers via Docker Compose. Em seguida, valide a saúde da API em http://127.0.0.1:8000/health e confirme que o Painel Streamlit em http://localhost:8501 está rodando perfeitamente.\"")
    ]

    for p_title, p_body in prompts:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(4)
        p_item.paragraph_format.space_after = Pt(2)
        r_t = p_item.add_run(f"🔹 {p_title}\n")
        r_t.bold = True
        r_t.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
        r_b = p_item.add_run(p_body)
        r_b.italic = True
        r_b.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # -------------------------------------------------------------
    # SEÇÃO 5: INICIALIZAÇÃO DO SISTEMA (DOCKER / LOCAL)
    # -------------------------------------------------------------
    add_heading_styled(doc, "5. Como Iniciar o Sistema na Nova Escola", level=1)

    add_heading_styled(doc, "Método A: Inicialização via Docker (Recomendado)", level=2)
    p = doc.add_paragraph()
    p.add_run("Abra o terminal na pasta do projeto e execute:\n")
    p.add_run("docker-compose up -d --build\n").bold = True
    p.add_run("Os serviços estarão prontos nos endereços:\n")
    p.add_run("• Painel de Gestão: ").bold = True
    p.add_run("http://localhost:8501\n")
    p.add_run("• Documentação API: ").bold = True
    p.add_run("http://localhost:8000/docs\n")
    p.add_run("• Orquestrador n8n: ").bold = True
    p.add_run("http://localhost:5678\n")

    add_heading_styled(doc, "Método B: Inicialização Local com Python Nativo", level=2)
    p = doc.add_paragraph()
    p.add_run("1. Crie e ative o ambiente virtual:\n")
    p.add_run("   python -m venv .venv\n   .\\.venv\\Scripts\\Activate.ps1\n")
    p.add_run("2. Instale os requisitos:\n")
    p.add_run("   pip install -r requirements.txt\n")
    p.add_run("3. Em um terminal, inicie a API:\n")
    p.add_run("   uvicorn app.main:app --host 0.0.0.0 --port 8000 --reload\n")
    p.add_run("4. Em outro terminal, inicie o Painel:\n")
    p.add_run("   streamlit run painel.py\n")

    # -------------------------------------------------------------
    # SEÇÃO 6: ROTINA DIÁRIA & TROUBLESHOOTING
    # -------------------------------------------------------------
    add_heading_styled(doc, "6. Rotina de Operação e Resolução de Problemas", level=1)

    p = doc.add_paragraph()
    p.add_run("1. ")
    p.add_run("07h30 / 13h30: ").bold = True
    p.add_run("Secretaria lança ausências e abre o Painel PAI (http://localhost:8501).\n")
    p.add_run("2. ")
    p.add_run("Disparo Automático: ").bold = True
    p.add_run("Clica em 'Disparar Notificações via WhatsApp' para alertar os responsáveis.\n")
    p.add_run("3. ")
    p.add_run("Triagem com IA: ").bold = True
    p.add_run("A IA atende os pais, registra os motivos de faltas e responde dúvidas.\n")
    p.add_run("4. ")
    p.add_run("Relatórios: ").bold = True
    p.add_run("Exportação de relatórios Word/PDF para prestação de contas à Diretoria de Ensino.\n")

    # Salvar documento
    output_path = Path("Manual_Replicacao_Presenca_Ativa_Inteligente.docx")
    doc.save(str(output_path))
    print(f"Documento salvo com sucesso em: {output_path.resolve()}")

if __name__ == "__main__":
    build_tutorial_docx()
